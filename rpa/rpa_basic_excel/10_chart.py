from openpyxl import load_workbook
wb = load_workbook("james.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, Reference, LineChart

# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3) # 범위 설정
# bar_chart= BarChart() # 차트 종류 설정 (line chart 등등 있음)
# bar_chart.add_data(bar_value) # 생성된 차트에 데이터 입력
# ws.add_chart(bar_chart, "E1") #차트를 넣을 위치

line_value = Reference(ws, min_row = 1, max_row = 11, min_col = 2, max_col = 3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data = True) #  계열 부분을 제목(영어, 수학) 에서 가져옴
line_chart.title = "성적표"
line_chart.style = 10
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"
ws.add_chart(line_chart, "E1")



wb.save("sample_chart.xlsx")